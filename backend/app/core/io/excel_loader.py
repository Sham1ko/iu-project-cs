from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict, Iterable, List, Sequence

from io import BytesIO

from openpyxl import load_workbook


_SHEET_ALIASES = {
    "subjects": ("subjects", "subject"),
    "teachers": ("teachers", "teacher"),
    "classes": ("classes", "class"),
    "hours": ("hours", "hour"),
}


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().lower()


def _iter_rows(ws) -> List[Dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [_normalize_header(value) for value in rows[0]]
    output: List[Dict[str, Any]] = []
    for raw in rows[1:]:
        if raw is None:
            continue
        if all(cell is None or str(cell).strip() == "" for cell in raw):
            continue
        row: Dict[str, Any] = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            value = raw[idx] if idx < len(raw) else None
            row[header] = value
        output.append(row)
    return output


def _find_sheet(workbook, kind: str):
    aliases = _SHEET_ALIASES[kind]
    for sheet_name in workbook.sheetnames:
        normalized = sheet_name.strip().lower()
        if normalized in aliases:
            return workbook[sheet_name]
    return None


def _find_sheet_by_name(workbook, target: str):
    target_norm = target.strip().lower()
    for sheet_name in workbook.sheetnames:
        if sheet_name.strip().lower() == target_norm:
            return workbook[sheet_name]
    return None


def _ensure_columns(rows: Iterable[Dict[str, Any]], required: Iterable[str], sheet: str) -> None:
    required_set = set(required)
    if not rows:
        raise ValueError(f"Excel sheet '{sheet}' is empty or missing required rows.")
    available = set(rows[0].keys())
    missing = [col for col in required_set if col not in available]
    if missing:
        raise ValueError(
            f"Excel sheet '{sheet}' is missing required columns: {', '.join(missing)}."
        )


def _parse_optional_int(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip()
    if text == "":
        return None
    return int(float(text))


def _normalize_class_name(name: Any) -> str:
    if name is None:
        return ""
    text = str(name).strip()
    text = re.sub(r"\s+", "", text)
    return text.lower()


def _extract_grade_from_class_name(class_name: str) -> int:
    match = re.match(r"^(\d+)", class_name.strip())
    if not match:
        raise ValueError(
            f"Class '{class_name}' must start with a numeric grade (e.g., 9A)."
        )
    return int(match.group(1))


def _tokenize_list_value(value: Any) -> List[str]:
    if value is None:
        return []
    if isinstance(value, bool):
        return []

    if isinstance(value, int):
        return [str(value)]

    if isinstance(value, float):
        if value.is_integer():
            return [str(int(value))]
        text = format(value, "g")
    else:
        text = str(value).strip()

    if text == "":
        return []

    if re.search(r"[;,]", text):
        return [part.strip() for part in re.split(r"[;,]", text) if part.strip()]

    if re.fullmatch(r"\d+(?:\.\d+)+", text):
        return [part.strip() for part in text.split(".") if part.strip()]

    tokens = text.split()
    if len(tokens) > 1 and all(token.isdigit() for token in tokens):
        return [token.strip() for token in tokens if token.strip()]

    return [text]


def _parse_subjects(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    _ensure_columns(rows, ["name"], "subjects")
    subjects: List[Dict[str, Any]] = []
    next_id = 1
    for row in rows:
        name = row.get("name")
        if name is None or str(name).strip() == "":
            raise ValueError("Subjects sheet requires a non-empty 'name' value.")
        subject_id = _parse_optional_int(row.get("id"))
        if subject_id is None:
            subject_id = next_id
        next_id = max(next_id, subject_id + 1)
        subject: Dict[str, Any] = {"id": subject_id, "name": str(name).strip()}
        for key in ("weekly_hours", "hours_per_week", "hours"):
            if key in row:
                subject["weekly_hours"] = _parse_optional_int(row.get(key))
                break
        subjects.append(subject)
    return subjects


def _parse_classes(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    _ensure_columns(rows, ["name", "grade"], "classes")
    classes: List[Dict[str, Any]] = []
    next_id = 1
    for row in rows:
        name = row.get("name")
        grade = row.get("grade")
        if name is None or str(name).strip() == "":
            raise ValueError("Classes sheet requires a non-empty 'name' value.")
        if grade is None or str(grade).strip() == "":
            raise ValueError("Classes sheet requires a non-empty 'grade' value.")
        class_id = _parse_optional_int(row.get("id"))
        if class_id is None:
            class_id = next_id
        next_id = max(next_id, class_id + 1)
        classes.append(
            {
                "id": class_id,
                "name": str(name).strip(),
                "grade": int(float(grade)),
            }
        )
    return classes


def _parse_subject_list(value: Any, subject_name_to_id: Dict[str, int]) -> List[int]:
    parts = _tokenize_list_value(value)
    subject_ids: List[int] = []

    for part in parts:
        cleaned = part.strip()
        if cleaned == "":
            continue
        if cleaned.isdigit():
            subject_ids.append(int(cleaned))
            continue
        key = cleaned.lower()
        if key not in subject_name_to_id:
            raise ValueError(f"Unknown subject '{cleaned}' in teachers sheet.")
        subject_ids.append(subject_name_to_id[key])

    # Preserve order while removing duplicates.
    seen: set[int] = set()
    unique_ids: List[int] = []
    for subject_id in subject_ids:
        if subject_id in seen:
            continue
        seen.add(subject_id)
        unique_ids.append(subject_id)
    return unique_ids


def _parse_class_list(
    value: Any,
    *,
    class_name_to_id: Dict[str, int],
    class_ids: set[int],
) -> List[int]:
    parts = _tokenize_list_value(value)
    groups: List[int] = []
    for part in parts:
        cleaned = part.strip()
        if cleaned == "":
            continue
        if cleaned.isdigit():
            class_id = int(cleaned)
            if class_id not in class_ids:
                raise ValueError(
                    f"Unknown class id '{class_id}' in teachers sheet groups."
                )
            groups.append(class_id)
            continue

        normalized = _normalize_class_name(cleaned)
        if normalized not in class_name_to_id:
            raise ValueError(f"Unknown class '{cleaned}' in teachers sheet groups.")
        groups.append(class_name_to_id[normalized])

    # Preserve order while removing duplicates.
    seen: set[int] = set()
    unique_groups: List[int] = []
    for group_id in groups:
        if group_id in seen:
            continue
        seen.add(group_id)
        unique_groups.append(group_id)
    return unique_groups


def _parse_teachers(
    rows: List[Dict[str, Any]],
    subject_name_to_id: Dict[str, int],
    *,
    class_name_to_id: Dict[str, int] | None = None,
    class_ids: set[int] | None = None,
) -> List[Dict[str, Any]]:
    _ensure_columns(rows, ["name", "subjects"], "teachers")
    teachers: List[Dict[str, Any]] = []
    next_id = 1
    for row in rows:
        name = row.get("name")
        if name is None or str(name).strip() == "":
            raise ValueError("Teachers sheet requires a non-empty 'name' value.")
        teacher_id = _parse_optional_int(row.get("id"))
        if teacher_id is None:
            teacher_id = next_id
        next_id = max(next_id, teacher_id + 1)
        subjects = _parse_subject_list(row.get("subjects"), subject_name_to_id)
        teacher: Dict[str, Any] = {
            "id": teacher_id,
            "name": str(name).strip(),
            "subjects": subjects,
        }

        if class_name_to_id is not None and class_ids is not None:
            for key in ("groups", "group", "classes"):
                if key in row:
                    teacher["groups"] = _parse_class_list(
                        row.get(key),
                        class_name_to_id=class_name_to_id,
                        class_ids=class_ids,
                    )
                    break

        for key in ("max_weekly_hours", "max_hours", "weekly_hours"):
            if key in row:
                teacher["max_weekly_hours"] = _parse_optional_int(row.get(key))
                break
        teachers.append(teacher)
    return teachers


def _parse_hours_sheet_by_grade(hours_ws) -> tuple[List[Dict[str, Any]], List[int]]:
    rows = list(hours_ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel sheet 'hours' is empty.")

    header = list(rows[0])
    grade_headers: List[int] = []
    for cell in header[1:]:
        grade = _parse_optional_int(cell)
        if grade is None:
            continue
        grade_headers.append(grade)

    if not grade_headers:
        raise ValueError("Excel sheet 'hours' must define grade headers in row 1.")

    subjects: List[Dict[str, Any]] = []
    next_id = 1
    for row in rows[1:]:
        if not row or (row[0] is None or str(row[0]).strip() == ""):
            continue
        name = str(row[0]).strip()
        hours_by_grade: Dict[int, int] = {}
        for idx, grade in enumerate(grade_headers, start=1):
            value = row[idx] if idx < len(row) else None
            hours = _parse_optional_int(value)
            if hours is None:
                continue
            hours_by_grade[int(grade)] = int(hours)
        subject = {"id": next_id, "name": name, "weekly_hours_by_grade": hours_by_grade}
        if hours_by_grade and len(set(hours_by_grade.values())) == 1:
            subject["weekly_hours"] = next(iter(hours_by_grade.values()))
        subjects.append(subject)
        next_id += 1
    return subjects, grade_headers


def _find_subject_name_column(header: Sequence[Any]) -> int:
    normalized = [_normalize_header(value) for value in header]
    for idx, name in enumerate(normalized):
        if name in ("subject", "subject name", "subject_name"):
            return idx
    for idx, name in enumerate(normalized):
        if name == "name":
            return idx
    return 1 if len(header) > 1 else 0


def _find_subject_id_column(header: Sequence[Any]) -> int | None:
    normalized = [_normalize_header(value) for value in header]
    for idx, name in enumerate(normalized):
        if name == "id":
            return idx
    return None


def _parse_hours_sheet_by_class(
    hours_ws,
) -> tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    rows = list(hours_ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel sheet 'hours' is empty.")

    header = list(rows[0])
    if not header:
        raise ValueError("Excel sheet 'hours' must define a header row.")

    subject_col = _find_subject_name_column(header)
    id_col = _find_subject_id_column(header)

    class_columns: List[tuple[int, str]] = []
    for idx, value in enumerate(header):
        if idx == subject_col:
            continue
        if id_col is not None and idx == id_col:
            continue
        if value is None or str(value).strip() == "":
            continue
        class_name = str(value).strip()
        class_columns.append((idx, class_name))

    if not class_columns:
        raise ValueError(
            "Excel sheet 'hours' must define class columns (e.g., 9A, 10B)."
        )

    classes: List[Dict[str, Any]] = []
    class_by_name_norm: Dict[str, int] = {}
    for col_index, class_name in class_columns:
        class_norm = _normalize_class_name(class_name)
        if class_norm in class_by_name_norm:
            continue
        grade = _extract_grade_from_class_name(class_name)
        class_id = len(classes) + 1
        class_by_name_norm[class_norm] = class_id
        classes.append(
            {
                "id": class_id,
                "name": class_name,
                "grade": grade,
            }
        )

    class_by_id = {cls["id"]: cls for cls in classes}
    class_id_by_col = {
        col_index: class_by_name_norm[_normalize_class_name(class_name)]
        for col_index, class_name in class_columns
    }

    subjects: List[Dict[str, Any]] = []
    next_id = 1
    for row in rows[1:]:
        if not row:
            continue
        name_value = row[subject_col] if subject_col < len(row) else None
        if name_value is None or str(name_value).strip() == "":
            continue

        subject_name = str(name_value).strip()
        subject_id = None
        if id_col is not None and id_col < len(row):
            subject_id = _parse_optional_int(row[id_col])
        if subject_id is None:
            subject_id = next_id
        next_id = max(next_id, subject_id + 1)

        hours_by_class: Dict[int, int] = {}
        for col_index, class_id in class_id_by_col.items():
            value = row[col_index] if col_index < len(row) else None
            hours = _parse_optional_int(value)
            if hours is None:
                continue
            hours_by_class[class_id] = int(hours)

        weekly_hours_by_grade: Dict[int, int] = {}
        for class_id, hours in hours_by_class.items():
            grade = class_by_id[class_id]["grade"]
            if grade in weekly_hours_by_grade and weekly_hours_by_grade[grade] != hours:
                continue
            weekly_hours_by_grade[grade] = hours

        subject: Dict[str, Any] = {
            "id": subject_id,
            "name": subject_name,
            "weekly_hours_by_class": hours_by_class,
        }
        if weekly_hours_by_grade:
            subject["weekly_hours_by_grade"] = weekly_hours_by_grade
        if hours_by_class and len(set(hours_by_class.values())) == 1:
            subject["weekly_hours"] = next(iter(hours_by_class.values()))

        subjects.append(subject)

    return subjects, classes


def _hours_header_looks_like_class_layout(header: Sequence[Any]) -> bool:
    normalized = [_normalize_header(value) for value in header]
    if any(name in ("subject", "subject name", "subject_name") for name in normalized):
        return True
    for value in header[1:]:
        if value is None:
            continue
        text = str(value).strip()
        if re.match(r"^\d+[A-Za-z]+$", text):
            return True
    return False


def _parse_matrix_format(workbook) -> Dict[str, List[Dict[str, Any]]]:
    hours_ws = _find_sheet(workbook, "hours")
    if hours_ws is None:
        raise ValueError("Excel file does not contain required 'hours' sheet.")

    raw_rows = list(hours_ws.iter_rows(values_only=True))
    if not raw_rows:
        raise ValueError("Excel sheet 'hours' is empty.")

    header = list(raw_rows[0])
    if _hours_header_looks_like_class_layout(header):
        subjects, classes = _parse_hours_sheet_by_class(hours_ws)
        teachers_ws = _find_sheet(workbook, "teachers")
        if teachers_ws is None:
            raise ValueError(
                "Excel class-matrix format requires a 'teachers' sheet."
            )
        teachers_rows = _iter_rows(teachers_ws)
        subject_name_to_id = {s["name"].lower(): s["id"] for s in subjects}
        class_name_to_id = {
            _normalize_class_name(cls["name"]): cls["id"] for cls in classes
        }
        class_ids = {cls["id"] for cls in classes}
        teachers = _parse_teachers(
            teachers_rows,
            subject_name_to_id,
            class_name_to_id=class_name_to_id,
            class_ids=class_ids,
        )
        return {
            "subjects": subjects,
            "classes": classes,
            "teachers": teachers,
        }

    subjects, _ = _parse_hours_sheet_by_grade(hours_ws)
    subject_name_to_id = {s["name"].lower(): s["id"] for s in subjects}

    classes_map: Dict[str, Dict[str, Any]] = {}
    teachers_map: Dict[str, Dict[str, Any]] = {}

    for subject in subjects:
        subject_name = subject["name"]
        subject_ws = _find_sheet_by_name(workbook, subject_name)
        if subject_ws is None:
            continue

        rows = list(subject_ws.iter_rows(values_only=True))
        if not rows:
            continue

        subject_header = rows[0]
        for cell in subject_header[1:]:
            if cell is None or str(cell).strip() == "":
                continue
            class_name = str(cell).strip()
            if class_name not in classes_map:
                grade = _extract_grade_from_class_name(class_name)
                classes_map[class_name] = {
                    "id": len(classes_map) + 1,
                    "name": class_name,
                    "grade": grade,
                }

        for row in rows[1:]:
            if not row:
                continue
            teacher_name = row[0]
            if teacher_name is None or str(teacher_name).strip() == "":
                continue
            teacher_name = str(teacher_name).strip()
            if teacher_name not in teachers_map:
                teachers_map[teacher_name] = {
                    "id": len(teachers_map) + 1,
                    "name": teacher_name,
                    "subjects": [],
                }
            teacher_entry = teachers_map[teacher_name]
            subject_id = subject_name_to_id[subject_name.lower()]
            if subject_id not in teacher_entry["subjects"]:
                teacher_entry["subjects"].append(subject_id)

    return {
        "subjects": subjects,
        "classes": list(classes_map.values()),
        "teachers": list(teachers_map.values()),
    }


def _load_dataset_from_workbook(workbook) -> Dict[str, List[Dict[str, Any]]]:
    subjects_ws = _find_sheet(workbook, "subjects")
    classes_ws = _find_sheet(workbook, "classes")
    teachers_ws = _find_sheet(workbook, "teachers")

    if subjects_ws and classes_ws and teachers_ws:
        subjects_rows = _iter_rows(subjects_ws)
        classes_rows = _iter_rows(classes_ws)
        teachers_rows = _iter_rows(teachers_ws)

        subjects = _parse_subjects(subjects_rows)
        classes = _parse_classes(classes_rows)
        subject_name_to_id = {s["name"].lower(): s["id"] for s in subjects}
        teachers = _parse_teachers(teachers_rows, subject_name_to_id)

        return {"subjects": subjects, "classes": classes, "teachers": teachers}

    hours_ws = _find_sheet(workbook, "hours")
    if hours_ws is not None:
        return _parse_matrix_format(workbook)

    raise ValueError(
        "Excel file must contain either sheets named: subjects, classes, teachers "
        "or the matrix format with an 'hours' sheet and subject sheets."
    )


def load_dataset_from_excel(path: Path) -> Dict[str, List[Dict[str, Any]]]:
    workbook = load_workbook(filename=path, data_only=True)
    return _load_dataset_from_workbook(workbook)


def load_dataset_from_excel_bytes(data: bytes) -> Dict[str, List[Dict[str, Any]]]:
    workbook = load_workbook(filename=BytesIO(data), data_only=True)
    return _load_dataset_from_workbook(workbook)

    # handled in _load_dataset_from_workbook
