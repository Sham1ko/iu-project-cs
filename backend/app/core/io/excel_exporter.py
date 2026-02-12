from __future__ import annotations

from io import BytesIO
from typing import Any, Dict, Iterable, List

from openpyxl import Workbook


def _coerce_int(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    text = str(value).strip()
    if text == "":
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _normalize_class_name(value: Any) -> str:
    text = "" if value is None else str(value)
    return "".join(text.strip().split()).lower()


def _coerce_weekly_hours(subject: Dict[str, Any]) -> int | None:
    hours = subject.get("weekly_hours")
    if hours is not None:
        return int(hours)
    hours_by_grade = subject.get("weekly_hours_by_grade")
    if isinstance(hours_by_grade, dict) and hours_by_grade:
        values = {int(v) for v in hours_by_grade.values() if v is not None}
        if len(values) == 1:
            return next(iter(values))
    return None


def _join_subjects(subjects: Iterable[Any]) -> str:
    return ",".join(str(item) for item in subjects if item is not None)


def _should_export_matrix_format(
    subjects: List[Dict[str, Any]],
    teachers: List[Dict[str, Any]],
) -> bool:
    has_class_hours = any(
        isinstance(subject.get("weekly_hours_by_class"), dict)
        and bool(subject.get("weekly_hours_by_class"))
        for subject in subjects
    )
    has_teacher_groups = any(
        isinstance(teacher.get("groups"), list) and bool(teacher.get("groups"))
        for teacher in teachers
    )
    return has_class_hours or has_teacher_groups


def _resolve_subject_hours_for_class(
    subject: Dict[str, Any],
    *,
    class_id: int,
    class_name: str,
) -> int | None:
    by_class = subject.get("weekly_hours_by_class")
    if isinstance(by_class, dict):
        for class_key, raw_hours in by_class.items():
            key_as_id = _coerce_int(class_key)
            if key_as_id is not None and key_as_id == class_id:
                hours = _coerce_int(raw_hours)
                return hours if hours is not None else 0
            key_as_name = _normalize_class_name(class_key)
            if key_as_name == _normalize_class_name(class_name):
                hours = _coerce_int(raw_hours)
                return hours if hours is not None else 0
        return 0

    by_grade = subject.get("weekly_hours_by_grade")
    if isinstance(by_grade, dict):
        return None

    return _coerce_weekly_hours(subject)


def _group_list_to_names(
    groups: List[Any],
    *,
    class_by_id: Dict[int, str],
) -> str:
    names: List[str] = []
    for group in groups:
        class_id = _coerce_int(group)
        if class_id is not None and class_id in class_by_id:
            names.append(class_by_id[class_id])
            continue
        text = str(group).strip()
        if text:
            names.append(text)
    return ",".join(names)


def export_dataset_to_excel_bytes(payload: Dict[str, Any]) -> bytes:
    subjects: List[Dict[str, Any]] = payload.get("subjects") or []
    teachers: List[Dict[str, Any]] = payload.get("teachers") or []
    classes: List[Dict[str, Any]] = payload.get("classes") or []

    wb = Workbook()

    if _should_export_matrix_format(subjects, teachers):
        sorted_classes = sorted(
            classes,
            key=lambda cls: (
                _coerce_int(cls.get("grade")) or 0,
                str(cls.get("name") or ""),
            ),
        )
        class_names = [str(cls.get("name") or "") for cls in sorted_classes]

        ws_hours = wb.active
        ws_hours.title = "hours"
        ws_hours.append(["ID", "Subject name", *class_names])
        for subject in subjects:
            subject_id = _coerce_int(subject.get("id"))
            row = [subject_id, subject.get("name")]
            for cls in sorted_classes:
                class_id = _coerce_int(cls.get("id"))
                class_name = str(cls.get("name") or "")
                value = None
                if class_id is not None:
                    value = _resolve_subject_hours_for_class(
                        subject, class_id=class_id, class_name=class_name
                    )
                row.append(value)
            ws_hours.append(row)

        class_by_id = {
            _coerce_int(cls.get("id")): str(cls.get("name") or "")
            for cls in sorted_classes
            if _coerce_int(cls.get("id")) is not None
        }

        ws_teachers = wb.create_sheet("teachers")
        ws_teachers.append(["ID", "name", "subjects", "groups", "max_hours"])
        for teacher in teachers:
            groups = teacher.get("groups")
            group_cell = ""
            if isinstance(groups, list) and groups:
                group_cell = _group_list_to_names(groups, class_by_id=class_by_id)
            ws_teachers.append(
                [
                    _coerce_int(teacher.get("id")),
                    teacher.get("name"),
                    _join_subjects(teacher.get("subjects", [])),
                    group_cell,
                    teacher.get("max_weekly_hours"),
                ]
            )

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    ws_subjects = wb.active
    ws_subjects.title = "subjects"
    ws_subjects.append(["id", "name", "weekly_hours"])
    for subject in subjects:
        ws_subjects.append(
            [
                subject.get("id"),
                subject.get("name"),
                _coerce_weekly_hours(subject),
            ]
        )

    ws_teachers = wb.create_sheet("teachers")
    ws_teachers.append(["id", "name", "subjects", "max_weekly_hours"])
    for teacher in teachers:
        ws_teachers.append(
            [
                teacher.get("id"),
                teacher.get("name"),
                _join_subjects(teacher.get("subjects", [])),
                teacher.get("max_weekly_hours"),
            ]
        )

    ws_classes = wb.create_sheet("classes")
    ws_classes.append(["id", "name", "grade"])
    for cls in classes:
        ws_classes.append([cls.get("id"), cls.get("name"), cls.get("grade")])

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
