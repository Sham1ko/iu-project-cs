from pathlib import Path

from openpyxl import Workbook

from app.core.io.excel_loader import load_dataset_from_excel


def _write_workbook(path: Path) -> None:
    wb = Workbook()

    subjects = wb.active
    subjects.title = "subjects"
    subjects.append(["id", "name", "weekly_hours"])
    subjects.append([1, "Math", 3])
    subjects.append([2, "English", 4])

    teachers = wb.create_sheet("teachers")
    teachers.append(["id", "name", "subjects", "max_weekly_hours"])
    teachers.append([1, "Alice", "1,2", 18])
    teachers.append([2, "Bob", "Math", 16])

    classes = wb.create_sheet("classes")
    classes.append(["id", "name", "grade"])
    classes.append([1, "5A", 5])

    wb.save(path)


def test_excel_loader_parses_dataset(tmp_path: Path) -> None:
    file_path = tmp_path / "dataset.xlsx"
    _write_workbook(file_path)

    dataset = load_dataset_from_excel(file_path)

    assert dataset["subjects"] == [
        {"id": 1, "name": "Math", "weekly_hours": 3},
        {"id": 2, "name": "English", "weekly_hours": 4},
    ]
    assert dataset["classes"] == [{"id": 1, "name": "5A", "grade": 5}]

    teachers = dataset["teachers"]
    assert teachers[0]["name"] == "Alice"
    assert teachers[0]["subjects"] == [1, 2]
    assert teachers[0]["max_weekly_hours"] == 18
    assert teachers[1]["name"] == "Bob"
    assert teachers[1]["subjects"] == [1]
    assert teachers[1]["max_weekly_hours"] == 16


def test_excel_loader_parses_hours_teachers_matrix_with_groups(tmp_path: Path) -> None:
    file_path = tmp_path / "dataset_matrix.xlsx"
    wb = Workbook()

    hours = wb.active
    hours.title = "hours"
    hours.append(["ID", "Subject name", "9A", "9B"])
    hours.append([1, "Math", 4, 4])
    hours.append([2, "English", 3, 3])

    teachers = wb.create_sheet("teachers")
    teachers.append(["ID", "name", "subjects", "groups", "max_hours"])
    teachers.append([1, "Alice", 1.2, "9A", 20])
    teachers.append([2, "Bob", 1.2, "9B", 20])

    wb.save(file_path)

    dataset = load_dataset_from_excel(file_path)

    assert dataset["classes"] == [
        {"id": 1, "name": "9A", "grade": 9},
        {"id": 2, "name": "9B", "grade": 9},
    ]
    assert dataset["subjects"][0]["id"] == 1
    assert dataset["subjects"][0]["name"] == "Math"
    assert dataset["subjects"][0]["weekly_hours_by_class"] == {1: 4, 2: 4}
    assert dataset["subjects"][1]["id"] == 2
    assert dataset["subjects"][1]["name"] == "English"
    assert dataset["subjects"][1]["weekly_hours_by_class"] == {1: 3, 2: 3}

    parsed_teachers = dataset["teachers"]
    assert parsed_teachers[0]["name"] == "Alice"
    assert parsed_teachers[0]["subjects"] == [1, 2]
    assert parsed_teachers[0]["groups"] == [1]
    assert parsed_teachers[0]["max_weekly_hours"] == 20
    assert parsed_teachers[1]["name"] == "Bob"
    assert parsed_teachers[1]["subjects"] == [1, 2]
    assert parsed_teachers[1]["groups"] == [2]
    assert parsed_teachers[1]["max_weekly_hours"] == 20
